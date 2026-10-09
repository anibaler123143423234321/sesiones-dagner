# -*- coding: utf-8 -*-
"""Examen final: la planilla de calificación del docente, con los estudiantes de la lista de asistencia.
Parte A (cuestionario, 0 a 10) + Parte B (presentación: cinco criterios de 0 a 2) = EF (0 a 20).
uso: python3 planilla_ef.py <repo>   → <repo>/SESION 16/docente/Calificacion_EF.xlsx
"""
import glob
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

REPO = sys.argv[1] if len(sys.argv) > 1 else '/home/user/sesiones-dagner'
SALIDA = os.path.join(REPO, 'SESION 16', 'docente', 'Calificacion_EF.xlsx')


def estudiantes():
    """Los nombres de la columna «APELLIDOS Y NOMBRES» de la primera hoja de asistencia que la tenga."""
    for ruta in glob.glob(os.path.join(REPO, 'ASISTENCIAS', '*.xlsx')):
        wb = load_workbook(ruta, data_only=True, read_only=True)
        for ws in wb.worksheets:
            filas = list(ws.iter_rows(values_only=True))
            for i, fila in enumerate(filas):
                for j, celda in enumerate(fila):
                    if isinstance(celda, str) and celda.strip().upper() == 'APELLIDOS Y NOMBRES':
                        nombres = []
                        for f in filas[i + 1:]:
                            v = f[j] if j < len(f) else None
                            if isinstance(v, str) and v.strip() and v.strip().upper() == v.strip() and len(v.split()) >= 3:
                                nombres.append(v.strip())
                            elif nombres:
                                break
                        if nombres:
                            return nombres
    return []


CRITERIOS = ['Sitio publicado y funcional', 'Diseño y fidelidad', 'Código y pruebas', 'Uso responsable de la IA', 'Presentación']
TEAL = PatternFill('solid', fgColor='177E89')
SUAVE = PatternFill('solid', fgColor='E6F8FA')
GRIS = PatternFill('solid', fgColor='F3F4F6')
borde = Border(*(Side(style='thin', color='CFD5D9'),) * 4)

wb = Workbook()
ws = wb.active
ws.title = 'EF'
ws['A1'] = 'Examen final [EF] · Diseño Web · Centro de Informática USS · Jueves 05/11/2026'
ws['A1'].font = Font(bold=True, size=13, color='0B0F19')
ws['A2'] = ('Parte A: nota del cuestionario del Aula Virtual (0 a 10). Parte B: cada criterio de 0 a 2 '
            '(logrado 2, en proceso 1, no logrado 0). El EF vale el 30 % del promedio final.')
ws['A2'].font = Font(italic=True, size=9, color='4B5563')
cab = ['N.º', 'Apellidos y nombres', 'Orden', 'Parte A (0-10)'] + CRITERIOS + ['Parte B (0-10)', 'EF (0-20)', 'Observaciones']
ws.append([])
ws.append(cab)
fila_cab = 4
for c, _ in enumerate(cab, start=1):
    celda = ws.cell(row=fila_cab, column=c)
    celda.font = Font(bold=True, color='FFFFFF', size=10)
    celda.fill = TEAL
    celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    celda.border = borde
ws.row_dimensions[fila_cab].height = 42

nombres = estudiantes()
nota_a = DataValidation(type='decimal', operator='between', formula1='0', formula2='10', allow_blank=True)
crit = DataValidation(type='whole', operator='between', formula1='0', formula2='2', allow_blank=True)
nota_a.error = 'La Parte A va de 0 a 10.'
crit.error = 'Cada criterio va de 0 a 2.'
ws.add_data_validation(nota_a)
ws.add_data_validation(crit)
for i, nombre in enumerate(nombres, start=1):
    r = fila_cab + i
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=nombre)
    ws.cell(row=r, column=10, value=f'=IF(COUNT(E{r}:I{r})=0,"",SUM(E{r}:I{r}))')
    ws.cell(row=r, column=11, value=f'=IF(AND(D{r}="",J{r}=""),"",N(D{r})+N(J{r}))')
    nota_a.add(f'D{r}')
    crit.add(f'E{r}:I{r}')
    for c in range(1, len(cab) + 1):
        celda = ws.cell(row=r, column=c)
        celda.border = borde
        celda.alignment = Alignment(horizontal='left' if c in (2, 12) else 'center', vertical='center')
        if c in (10, 11):
            celda.fill = SUAVE
            celda.font = Font(bold=True)
        elif i % 2 == 0:
            celda.fill = GRIS
ult = fila_cab + len(nombres)
ws.cell(row=ult + 2, column=2, value='Promedio del aula').font = Font(bold=True)
for col in 'DJK':
    c = ws[f'{col}{ult + 2}']
    c.value = f'=IFERROR(ROUND(AVERAGE({col}{fila_cab + 1}:{col}{ult}),2),"")'
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal='center')
anchos = {'A': 5, 'B': 40, 'C': 7, 'D': 11, 'E': 12, 'F': 12, 'G': 12, 'H': 12, 'I': 13, 'J': 11, 'K': 10, 'L': 34}
for col, w in anchos.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'C5'

# hoja 2: la pauta, para tenerla a mano al calificar
p = wb.create_sheet('Pauta Parte B')
p.append(['Criterio', 'Logrado (2)', 'En proceso (1)', 'No logrado (0)'])
p.append(['Sitio publicado y funcional', 'En Netlify; enlaces, celular a 390 px, formulario que guarda mensajes y 404', 'Publicado, con algún enlace roto o sin probar en el celular', 'No está publicado o no abre'])
p.append(['Diseño y fidelidad', 'Figma con escritorio, celular y componentes; el sitio respeta sus colores, letras y medidas', 'Hay Figma, pero el sitio se aparta en varias partes', 'Sin Figma o sin relación con el sitio'])
p.append(['Código y pruebas', 'HTML semántico y válido; explica la parte del código que se le pide; Lighthouse 90+ en Accesibilidad', 'Explica con ayuda o hay errores en el validador', 'No puede explicar su código'])
p.append(['Uso responsable de la IA', 'Bitácora con 3 prompts o más: qué pidió, qué corrigió y cómo lo comprobó', 'Bitácora incompleta o sin la comprobación', 'Sin bitácora, o código que no entiende'])
p.append(['Presentación', 'Sigue el guion en 5 minutos, con orden y claridad; responde las preguntas', 'Se pasa del tiempo o salta partes del guion', 'No presenta'])
for c in p[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = TEAL
for fila in p.iter_rows(min_row=1, max_row=6):
    for c in fila:
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = borde
for col, w in {'A': 26, 'B': 48, 'C': 36, 'D': 30}.items():
    p.column_dimensions[col].width = w

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
print('ok', SALIDA, len(nombres), 'estudiantes')
